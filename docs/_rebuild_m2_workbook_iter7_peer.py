"""Bangun ulang sheet M2aLowIrradiance (iterasi 7) -- detektor relatif-tetangga.

Detektor OLS lama dihapus 2026-09-28 (docs/M2_Family_Summary.md #6).
``_extend_m2_workbook_iter7.py`` tidak bisa dijalankan ulang (ia mengharapkan
29 sheet hasil iterasi 6; workbook kini 46), jadi skrip ini MENGGANTI isi
keempat sheet LI di posisinya semula dan menyesuaikan Config:
  baris 52-58 (nama sel tetap): pita POA, low_ratio_threshold (bekas
  slope_threshold), robust_z_min (bekas r_squared_min), min_low_samples = 12;
  3 baris baru di akhir Config: low_band_min_elevation_deg, min_peers,
  min_mid_samples.

Rantai formula (tanpa array formula, aman LibreOffice):
  Helpers_LI per timestamp: pr_k = P_k/POA; midpr_k = pr_k bila POA di pita
  menengah; e_k = pr_k / MEDIAN(midpr_k) (hanya bila n_mid_k >= min);
  peer = MEDIAN(e_1..e_6) bila COUNT >= min_peers; dev_k = e_k / peer;
  lowdev_k = dev_k bila POA di pita rendah DAN elevasi >= batas.
  M2a_LowIrradiance per inverter: low_ratio = MEDIAN(lowdev_k); z terhadap
  median/MAD inverter terevaluasi (MAD dibatasi 0,005, skala 1,4826).

Data demo SINTETIS (WB05, 6 inverter): faktor cuaca 0,93 untuk semua di
cahaya rendah (tidak di-flag), INV03 buruk di cahaya rendah (HIGH), INV05
buruk hanya saat matahari rendah (lolos gerbang elevasi), INV06 tanpa data
saat cahaya rendah (insufficient_data). Diverifikasi terhadap
M2aLowIrradiance asli + replika Python rantai formula (LibreOffice tidak
tersedia untuk recalc langsung).
"""
from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.workbook.defined_name import DefinedName

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))
from pv_pipeline.m2a.low_irradiance import M2aLowIrradiance  # noqa: E402

INPUT = Path(__file__).parent / "M2_PV_Performance_Workbook.xlsx"
LI_SHEETS = ["Raw_Data_LI", "Helpers_LI", "M2a_LowIrradiance", "LI_Summary"]

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=13, bold=True, color="305496")
NOTE_FONT = Font(italic=True, size=9, color="808080")
SEV_FILL = {
    "CRITICAL": PatternFill("solid", fgColor="E06666"),
    "HIGH": PatternFill("solid", fgColor="F6B26B"),
    "MEDIUM": PatternFill("solid", fgColor="FFE599"),
}
LOW_FILL = PatternFill("solid", fgColor="DDEBF7")

# Satu sumber parameter: Config sheet DAN detektor verifikasi.
CFG = {
    "poa_low_min": 50.0, "poa_low_max": 250.0, "poa_mid_min": 300.0, "poa_mid_max": 800.0,
    "low_ratio_threshold": 0.90, "robust_z_min": 3.0, "min_low_samples": 12,
    "low_band_min_elevation_deg": 30.0, "min_peers": 5, "min_mid_samples": 30,
}
MAD_FLOOR, MAD_SCALE = 0.005, 1.4826


def set_header(ws, row, headers, start=1):
    for ci, h in enumerate(headers, start=start):
        c = ws.cell(row=row, column=ci, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER


def title_note(ws, title, notes):
    ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    for i, nt in enumerate(notes, start=2):
        ws.cell(row=i, column=1, value=nt).font = NOTE_FONT


# ===========================================================================
# Data demo sintetis
# ===========================================================================
TS = pd.date_range("2026-05-14 06:30", "2026-05-14 17:30", freq="5min")
HRS = np.asarray(TS.hour + TS.minute / 60.0)
IDX = np.arange(len(TS))
CLOUD = (HRS >= 11) & (HRS < 13)
POA = np.where(CLOUD, 150.0 + 40.0 * np.sin(IDX), 1000.0 * np.sin(np.pi * (HRS - 6) / 12) ** 2)
ELEV = 85.0 * np.sin(np.pi * (HRS - 6) / 12)
LOW = POA < 250.0
INVERTERS = [f"WB05-INV{k:02d}" for k in range(1, 7)]


def demo_power(k: int) -> np.ndarray:
    f = np.where(LOW, 0.93, 1.0)                             # cuaca: seragam se-WB
    if k == 3:
        f = f * np.where(LOW, 0.83, 1.0)                     # cacat low-light
    if k == 5:
        f = f * np.where(LOW & (ELEV < 30.0), 0.80, 1.0)     # bayangan pagi/sore
    p = (100.0 + 10.0 * k) * POA / 1000.0 * f * (1.0 + 0.01 * np.sin(0.37 * IDX + k))
    if k == 6:
        p = np.where(LOW, np.nan, p)                         # tanpa data saat redup
    return p


P = {inv: demo_power(k) for k, inv in enumerate(INVERTERS, start=1)}
N = len(TS)
DF, DL = 7, 7 + N - 1

# ===========================================================================
# Verifikasi: replika Python rantai formula vs detektor asli
# ===========================================================================


def replica() -> pd.DataFrame:
    in_mid = (POA >= CFG["poa_mid_min"]) & (POA <= CFG["poa_mid_max"])
    in_low = ((POA >= CFG["poa_low_min"]) & (POA <= CFG["poa_low_max"])
              & (ELEV >= CFG["low_band_min_elevation_deg"]))
    pr = np.column_stack([P[inv] / POA for inv in INVERTERS])
    mid_ok = in_mid[:, None] & np.isfinite(pr) & (pr > 0)
    n_mid = mid_ok.sum(axis=0)
    mid_med = np.array([np.median(pr[mid_ok[:, j], j]) for j in range(len(INVERTERS))])
    e = np.where(n_mid >= CFG["min_mid_samples"], pr / mid_med, np.nan)
    n_peers = np.isfinite(e).sum(axis=1)
    peer = np.where(n_peers >= CFG["min_peers"], np.nanmedian(e, axis=1), np.nan)
    dev = e / peer[:, None]
    lowdev = np.where(in_low[:, None] & np.isfinite(dev), dev, np.nan)
    n_low = np.isfinite(lowdev).sum(axis=0)
    evaluated = (n_low >= CFG["min_low_samples"]) & (n_mid >= CFG["min_mid_samples"])
    low_ratio = np.array([np.nanmedian(lowdev[:, j]) if evaluated[j] else np.nan
                          for j in range(len(INVERTERS))])
    med = np.median(low_ratio[evaluated])
    mad = np.median(np.abs(low_ratio[evaluated] - med))
    z = (low_ratio - med) / (MAD_SCALE * max(mad, MAD_FLOOR))
    flagged = evaluated & (low_ratio < CFG["low_ratio_threshold"]) & (z < -CFG["robust_z_min"])
    cls = np.where(~evaluated, "insufficient_data",
                   np.where(flagged, "low_irradiance_underperform", "normal"))
    return pd.DataFrame({"inverter_id": INVERTERS, "n_low": n_low, "low_ratio": low_ratio,
                         "robust_z": np.where(evaluated, z, np.nan), "classification": cls})


class _DemoPOA:
    def _lookup(self, values, timestamps):
        idx = pd.DatetimeIndex(timestamps)
        return pd.Series(pd.Series(values, index=TS).reindex(idx).to_numpy(), index=idx)

    def get_poa(self, timestamps, wb_id, source="auto"):
        return self._lookup(POA, timestamps)

    def get_solar_elevation(self, timestamps):
        return self._lookup(ELEV, timestamps)


def detector() -> pd.DataFrame:
    rows = []
    for inv in INVERTERS:
        for t, p in zip(TS, P[inv]):
            if np.isfinite(p):
                rows.append({"Inverter_ID": inv, "Start Time": t,
                             "PV1 Power(kW)": p / 2, "PV2 Power(kW)": p / 2})
    cfg = {"m2a_low_irradiance": {
        "enabled": True,
        "poa_low_range": [CFG["poa_low_min"], CFG["poa_low_max"]],
        "poa_mid_range": [CFG["poa_mid_min"], CFG["poa_mid_max"]],
        **{k: CFG[k] for k in ("low_ratio_threshold", "robust_z_min", "min_low_samples",
                                "low_band_min_elevation_deg", "min_peers", "min_mid_samples")},
        "hour_range": [6.0, 18.0], "hour_cutoff_end": 18.0, "solar_elevation_min_deg": 5.0,
        "respect_inverter_shutdown": False, "pv_max": 2,
    }}
    sm = M2aLowIrradiance(poa=_DemoPOA())
    sm.run(pd.DataFrame(rows), cfg)
    return sm.artifacts["LowIrradianceFit"].set_index("inverter_id")


rep = replica().set_index("inverter_id")
det = detector()
assert (rep["classification"] == det["classification"]).all(), (rep, det)
assert (rep["n_low"] == det["n_low_samples"]).all()
ev = rep["classification"] != "insufficient_data"
assert np.allclose(rep.loc[ev, "low_ratio"], det.loc[ev, "low_ratio"], rtol=0, atol=1e-12)
assert np.allclose(rep.loc[ev, "robust_z"], det.loc[ev, "robust_z"], rtol=0, atol=1e-9)
expected = {"WB05-INV03": "low_irradiance_underperform", "WB05-INV06": "insufficient_data"}
for inv in INVERTERS:
    assert rep.loc[inv, "classification"] == expected.get(inv, "normal"), inv
assert 0.80 <= rep.loc["WB05-INV03", "low_ratio"] < 0.85, "INV03 harus HIGH"
print("Verifikasi replika == detektor:")
print(rep.round(4).to_string())

# ===========================================================================
# Workbook
# ===========================================================================
wb = load_workbook(INPUT)
assert all(s in wb.sheetnames for s in LI_SHEETS), "sheet LI tidak ditemukan"
before = list(wb.sheetnames)

# --- Config ----------------------------------------------------------------
ws = wb["Config"]
if "cfg_li_slope_threshold" in wb.defined_names:
    assert ws["B56"].value == "slope_threshold" and ws["B57"].value == "r_squared_min"
    ws["B56"], ws["C56"] = "low_ratio_threshold", CFG["low_ratio_threshold"]
    ws["D56"] = "low_ratio < ini -> kandidat flag (relatif median tetangga se-WB)"
    ws["B57"], ws["C57"] = "robust_z_min", CFG["robust_z_min"]
    ws["D57"] = "...DAN robust-z < -nilai ini terhadap inverter se-WB-hari"
    ws["C58"] = CFG["min_low_samples"]
    ws["D58"] = "min sampel pita rendah bermatahari tinggi (1 jam @5 menit)"
    del wb.defined_names["cfg_li_slope_threshold"]
    del wb.defined_names["cfg_li_r2_min"]
    wb.defined_names["cfg_li_low_ratio_threshold"] = DefinedName(
        "cfg_li_low_ratio_threshold", attr_text="Config!$C$56")
    wb.defined_names["cfg_li_robust_z_min"] = DefinedName(
        "cfg_li_robust_z_min", attr_text="Config!$C$57")
    last = max(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value is not None)
    assert last == 78, f"Config diharapkan berakhir di baris 78, dapat {last}"
    new_cfg = [
        ("low_band_min_elevation_deg", CFG["low_band_min_elevation_deg"],
         "pita rendah hanya saat matahari >= ini (awan siang, bukan matahari rendah)",
         "cfg_li_low_elev_min"),
        ("min_peers", CFG["min_peers"], "min inverter se-WB per timestamp untuk median tetangga",
         "cfg_li_min_peers"),
        ("min_mid_samples", CFG["min_mid_samples"], "min sampel pita menengah (normalisasi e)",
         "cfg_li_min_mid_samples"),
    ]
    for i, (key, val, note, name) in enumerate(new_cfg):
        ri = last + 1 + i
        ws.cell(row=ri, column=1, value="m2a_low_irradiance").border = BORDER
        ws.cell(row=ri, column=2, value=key).border = BORDER
        c = ws.cell(row=ri, column=3, value=val)
        c.border = BORDER
        c.fill = PatternFill("solid", fgColor="FFF2CC")
        ws.cell(row=ri, column=4, value=note).border = BORDER
        wb.defined_names[name] = DefinedName(name, attr_text=f"Config!$C${ri}")
    print("Config: baris 56-58 diganti, 79-81 ditambah.")
else:
    for name in ("cfg_li_low_ratio_threshold", "cfg_li_robust_z_min", "cfg_li_low_elev_min",
                 "cfg_li_min_peers", "cfg_li_min_mid_samples"):
        assert name in wb.defined_names, f"{name} hilang"
    print("Config sudah versi relatif-tetangga; dilewati.")

# --- README ----------------------------------------------------------------
readme = wb["README"]
assert readme["C12"].value == "M2aLowIrradiance"
readme["D12"] = ("Raw_Data_LI, Helpers_LI, M2a_LowIrradiance, LI_Summary "
                 "(dibangun ulang 2026-09-28: relatif-tetangga)")

# --- Ganti keempat sheet LI di posisi yang sama ---------------------------
sheets = {}
for name in LI_SHEETS:
    pos = wb.sheetnames.index(name)
    wb.remove(wb[name])
    sheets[name] = wb.create_sheet(name, pos)

# Raw_Data_LI
ws = sheets["Raw_Data_LI"]
title_note(ws, "Raw_Data_LI -- POA, elevasi matahari & daya inverter per timestamp (SINTETIS)", [
    "WB05, 6 inverter, satu hari @5 menit; awan 11:00-13:00 (POA ~150, matahari tinggi). "
    "P = sum daya per-PV (kW). Sel kosong = tidak ada data (INV06 saat POA < 250).",
    "Skenario: semua x0,93 di cahaya rendah (cuaca, tidak di-flag); INV03 x0,83 lagi (cacat low-light); "
    "INV05 x0,80 hanya saat elevasi < 30 derajat (bayangan pagi/sore).",
    "Ganti dengan data Huawei aktual satu WB (paste over): semua inverter WB berbagi POA WS yang sama.",
])
set_header(ws, 6, ["timestamp", "POA (W/m²)", "elev (°)"] + [f"{inv} P (kW)" for inv in INVERTERS])
for i in range(N):
    r = DF + i
    ws.cell(row=r, column=1, value=TS[i].to_pydatetime()).number_format = "yyyy-mm-dd hh:mm"
    ws.cell(row=r, column=2, value=float(POA[i])).number_format = "0.0"
    ws.cell(row=r, column=3, value=float(ELEV[i])).number_format = "0.0"
    for j, inv in enumerate(INVERTERS):
        v = P[inv][i]
        c = ws.cell(row=r, column=4 + j, value=float(v) if np.isfinite(v) else None)
        c.number_format = "0.000"
    for ci in range(1, 10):
        ws.cell(row=r, column=ci).border = BORDER
    if LOW[i]:
        ws.cell(row=r, column=2).fill = LOW_FILL
ws.column_dimensions["A"].width = 17

# Helpers_LI
ws = sheets["Helpers_LI"]
title_note(ws, "Helpers_LI -- efisiensi relatif & deviasi terhadap median tetangga se-WB", [
    "pr=P/POA; midpr=pr bila POA di pita menengah; e=pr/median(midpr) (baris 5) bila n_mid (baris 4) "
    ">= min; peer=MEDIAN(e) bila COUNT >= min_peers; dev=e/peer; lowdev=dev bila pita rendah & elevasi >= batas.",
    "Sel \"\" = tidak ada nilai; MEDIAN/COUNT mengabaikannya (sama dengan NaN di Python).",
])
PR, MID, E, DEV, LDEV = 6, 12, 18, 26, 32          # kolom awal tiap blok 6 inverter
NP, PEER = 24, 25
hdr = (["timestamp", "POA", "elev", "in_mid", "in_low_hs"]
       + [f"pr_{k}" for k in range(1, 7)] + [f"midpr_{k}" for k in range(1, 7)]
       + [f"e_{k}" for k in range(1, 7)] + ["n_peers", "peer_median"]
       + [f"dev_{k}" for k in range(1, 7)] + [f"lowdev_{k}" for k in range(1, 7)])
set_header(ws, 6, hdr)
ws.cell(row=4, column=MID - 1, value="n_mid").font = NOTE_FONT
ws.cell(row=5, column=MID - 1, value="median mid").font = NOTE_FONT
for j in range(6):
    mc = col(MID + j)
    ws.cell(row=4, column=MID + j, value=f"=COUNT({mc}{DF}:{mc}{DL})").border = BORDER
    c = ws.cell(row=5, column=MID + j, value=f"=MEDIAN({mc}{DF}:{mc}{DL})")
    c.border = BORDER
    c.number_format = "0.000000"
for i in range(N):
    r = DF + i
    ws.cell(row=r, column=1, value=f"=Raw_Data_LI!A{r}").number_format = "yyyy-mm-dd hh:mm"
    ws.cell(row=r, column=2, value=f"=Raw_Data_LI!B{r}").number_format = "0.0"
    ws.cell(row=r, column=3, value=f"=Raw_Data_LI!C{r}").number_format = "0.0"
    ws.cell(row=r, column=4, value=f"=IF(AND(B{r}>=cfg_li_poa_mid_min,B{r}<=cfg_li_poa_mid_max),1,0)")
    ws.cell(row=r, column=5, value=(
        f"=IF(AND(B{r}>=cfg_li_poa_low_min,B{r}<=cfg_li_poa_low_max,C{r}>=cfg_li_low_elev_min),1,0)"))
    for j in range(6):
        raw = f"Raw_Data_LI!{col(4 + j)}{r}"
        pr, mid, e = col(PR + j), col(MID + j), col(E + j)
        dev = col(DEV + j)
        ws.cell(row=r, column=PR + j, value=f'=IF({raw}="","",{raw}/B{r})').number_format = "0.000000"
        ws.cell(row=r, column=MID + j, value=(
            f'=IF(AND(D{r}=1,ISNUMBER({pr}{r}),{pr}{r}>0),{pr}{r},"")')).number_format = "0.000000"
        ws.cell(row=r, column=E + j, value=(
            f'=IF(AND(ISNUMBER({pr}{r}),{mid}$4>=cfg_li_min_mid_samples),{pr}{r}/{mid}$5,"")'
        )).number_format = "0.0000"
        ws.cell(row=r, column=DEV + j, value=(
            f'=IF(AND(ISNUMBER({e}{r}),ISNUMBER(Y{r})),{e}{r}/Y{r},"")')).number_format = "0.0000"
        ws.cell(row=r, column=LDEV + j, value=(
            f'=IF(AND(E{r}=1,ISNUMBER({dev}{r})),{dev}{r},"")')).number_format = "0.0000"
    e_rng = f"{col(E)}{r}:{col(E + 5)}{r}"
    ws.cell(row=r, column=NP, value=f"=COUNT({e_rng})")
    ws.cell(row=r, column=PEER, value=(
        f'=IF(X{r}>=cfg_li_min_peers,MEDIAN({e_rng}),"")')).number_format = "0.0000"
assert col(NP) == "X" and col(PEER) == "Y"
ws.column_dimensions["A"].width = 17

# M2a_LowIrradiance
ws = sheets["M2a_LowIrradiance"]
title_note(ws, "M2a_LowIrradiance -- low_ratio relatif tetangga se-WB + robust z + klasifikasi", [
    "low_ratio = MEDIAN(lowdev); median & MAD atas inverter terevaluasi; "
    f"z = (low_ratio - median) / ({MAD_SCALE} x MAX(MAD, {MAD_FLOOR})).",
    "Flag bila low_ratio < cfg_li_low_ratio_threshold DAN z < -cfg_li_robust_z_min. "
    "Severity: < 0,80 CRITICAL; < 0,85 HIGH; lainnya MEDIUM.",
])
set_header(ws, 4, ["inverter_id", "n_mid", "median_mid_pr", "n_low", "evaluated", "low_ratio",
                   "wb_median_low_ratio", "abs_dev", "mad", "robust_z", "classification",
                   "severity", "emit", "confidence", "message"])
R0, R1 = 5, 5 + len(INVERTERS) - 1
for j, inv in enumerate(INVERTERS):
    r = R0 + j
    mid, ldev = col(MID + j), col(LDEV + j)
    ws.cell(row=r, column=1, value=inv)
    ws.cell(row=r, column=2, value=f"=Helpers_LI!{mid}$4")
    ws.cell(row=r, column=3, value=f"=Helpers_LI!{mid}$5").number_format = "0.000000"
    ws.cell(row=r, column=4, value=f"=COUNT(Helpers_LI!{ldev}${DF}:{ldev}${DL})")
    ws.cell(row=r, column=5, value=(
        f"=IF(AND(D{r}>=cfg_li_min_low_samples,B{r}>=cfg_li_min_mid_samples),1,0)"))
    ws.cell(row=r, column=6, value=(
        f'=IF(E{r}=1,MEDIAN(Helpers_LI!{ldev}${DF}:{ldev}${DL}),"")')).number_format = "0.0000"
    ws.cell(row=r, column=7, value=f'=IF(E{r}=1,MEDIAN($F${R0}:$F${R1}),"")').number_format = "0.0000"
    ws.cell(row=r, column=8, value=f'=IF(E{r}=1,ABS(F{r}-G{r}),"")').number_format = "0.0000"
    ws.cell(row=r, column=9, value=f'=IF(E{r}=1,MEDIAN($H${R0}:$H${R1}),"")').number_format = "0.0000"
    ws.cell(row=r, column=10, value=(
        f'=IF(E{r}=1,(F{r}-G{r})/({MAD_SCALE}*MAX(I{r},{MAD_FLOOR})),"")')).number_format = "0.00"
    ws.cell(row=r, column=11, value=(
        f'=IF(E{r}=0,"insufficient_data",IF(AND(F{r}<cfg_li_low_ratio_threshold,'
        f'J{r}<-cfg_li_robust_z_min),"low_irradiance_underperform","normal"))'))
    ws.cell(row=r, column=12, value=(
        f'=IF(K{r}="low_irradiance_underperform",IF(F{r}<0.8,"CRITICAL",'
        f'IF(F{r}<0.85,"HIGH","MEDIUM")),"NORMAL")'))
    ws.cell(row=r, column=13, value=f'=IF(K{r}="low_irradiance_underperform",1,0)')
    ws.cell(row=r, column=14, value=f'=IF(M{r}=1,MIN(95,50+5*ABS(J{r})),"")').number_format = "0.0"
    ws.cell(row=r, column=15, value=(
        f'=IF(M{r}=1,"Low-light relatif tetangga: low_ratio="&TEXT(F{r},"0.000")&" < "'
        f'&cfg_li_low_ratio_threshold&" (median "&LEFT(A{r},4)&" "&TEXT(G{r},"0.000")'
        f'&", z="&TEXT(J{r},"0.0")&", n="&D{r}&")","")'))
    for ci in range(1, 16):
        ws.cell(row=r, column=ci).border = BORDER
for sev, fill in SEV_FILL.items():
    ws.conditional_formatting.add(f"L{R0}:L{R1}",
                                  CellIsRule(operator="equal", formula=[f'"{sev}"'], fill=fill))
ws.column_dimensions["A"].width = 13
ws.column_dimensions["K"].width = 26
ws.column_dimensions["O"].width = 60

# LI_Summary
ws = sheets["LI_Summary"]
title_note(ws, "LI_Summary -- hitung per klasifikasi (mirror LowIrradianceSummary)", [
    "general_underperform tidak lagi dihasilkan sejak 2026-09-28; baris dipertahankan (selalu 0) "
    "supaya dasbor lama tidak rusak.",
])
set_header(ws, 4, ["classification", "count"])
cats = ["low_irradiance_underperform", "general_underperform", "normal", "insufficient_data"]
for i, cat in enumerate(cats):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=cat).border = BORDER
    ws.cell(row=rr, column=2, value=(
        f'=COUNTIF(M2a_LowIrradiance!$K${R0}:$K${R1},"{cat}")')).border = BORDER
ws.cell(row=10, column=1, value="n_emit").border = BORDER
ws.cell(row=10, column=2, value=f"=SUM(M2a_LowIrradiance!M{R0}:M{R1})").border = BORDER
ws.column_dimensions["A"].width = 28

# ===========================================================================
# Simpan + periksa integritas
# ===========================================================================
wb.save(INPUT)
after = list(load_workbook(INPUT).sheetnames)
assert after == before, "urutan sheet berubah"
print(f"OK -- {len(after)} sheet; LI dibangun ulang di posisi "
      f"{[after.index(s) for s in LI_SHEETS]}.")
